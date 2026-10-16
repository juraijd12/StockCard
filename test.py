import openpyxl

FILE = r"D:\Desktop\CEILING WORKS.xlsx"

wb = openpyxl.load_workbook(FILE, data_only=False)
ws = wb["Sheet1"]

blocks = [
    (123, 183),
    (184, 242),
    (243, 301),
    (302, 361),
    (775, 833),
    (834, 890),
]

for start, end in blocks:
    print("\n" + "=" * 80)
    print(f"ROWS {start}-{end}")
    print("=" * 80)

    # Only inspect the metadata area around the header.
    for r in range(start, min(start + 7, end + 1)):
        values = []

        for c in range(1, 18):
            value = ws.cell(r, c).value

            if value is not None:
                col = ws.cell(r, c).column_letter
                values.append(f"{col}{r}={value!r}")

        if values:
            print(" | ".join(values))