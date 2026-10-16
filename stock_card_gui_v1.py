from __future__ import annotations

import sqlite3
import json
import os
import subprocess
import sys
import shutil
import hashlib
import tkinter as tk
import re
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from tkcalendar import DateEntry
try:
    from PIL import Image, ImageTk
except ImportError:
    Image = ImageTk = None
from stock_card_document_print import export_document

from stock_card_backend_v1 import DEFAULT_BATCH, DEFAULT_DB, add_paper, add_paper_record, add_paper_batch, add_material_alias, connect, create_document, create_material, post_document, register_project, set_material_description, set_transaction_status, update_paper_record


class StockCardApp(tk.Tk):
    def __init__(self, db_path: Path = DEFAULT_DB):
        super().__init__()
        self.title("AIYU Warehouse")
        self.geometry("1180x760")
        self.minsize(760, 480)
        self.configure(bg="#f3f6fb")
        self.db_path = db_path
        register_project(self.db_path, None)
        self.config_path = Path(__file__).resolve().parent / "settings.json"
        self.settings = self._load_settings()
        self._style()
        self.material_ids: list[int] = []
        self.selected_id: int | None = None
        self.selected_source: tuple[str, str, int, int, int] | None = None

        menu = tk.Menu(self)
        file_menu = tk.Menu(menu, tearoff=False)
        file_menu.add_command(label="Refresh inventory", command=self.refresh_inventory)
        file_menu.add_command(label="Export selected stock card", command=self.export_stock_card)
        file_menu.add_command(label="Export all stock cards", command=self.export_all_stock_cards)
        file_menu.add_separator()
        file_menu.add_command(label="Exit", command=self.destroy)
        view_menu = tk.Menu(menu, tearoff=False)
        view_menu.add_command(label="Show all materials", command=self.show_all)
        tools_menu = tk.Menu(menu, tearoff=False)
        tools_menu.add_command(label="Open Excel source", command=self.open_source)
        tools_menu.add_command(label="Settings", command=self.open_settings)
        help_menu = tk.Menu(menu, tearoff=False)
        help_menu.add_command(label="About AIYU Warehouse", command=lambda: messagebox.showinfo("About", "AIYU Warehouse\nStock-card and warehouse entry tool"))
        menu.add_cascade(label="File", menu=file_menu)
        menu.add_cascade(label="View", menu=view_menu)
        menu.add_cascade(label="Tools", menu=tools_menu)
        menu.add_cascade(label="Help", menu=help_menu)
        self.config(menu=menu)

        toolbar = ttk.Frame(self, style="Toolbar.TFrame", padding=(22, 7))
        toolbar.pack(fill="x", side="top")
        ttk.Button(toolbar, text="↻ Refresh inventory", command=self.refresh_inventory).pack(side="left", padx=(0, 6))
        ttk.Button(toolbar, text="Open Excel source", command=self.open_source).pack(side="left", padx=6)
        ttk.Button(toolbar, text="Export selected stock card", style="Accent.TButton", command=self.export_stock_card).pack(side="left", padx=6)
        ttk.Button(toolbar, text="Export all stock cards", command=self.export_all_stock_cards).pack(side="left", padx=6)
        ttk.Button(toolbar, text="＋ Add entry", style="Accent.TButton", command=self.open_entry_window).pack(side="left", padx=6)
        ttk.Button(toolbar, text="Receiving report", style="Accent.TButton", command=lambda: self.open_document_window(document_type="RECEIVING")).pack(side="left", padx=6)
        ttk.Button(toolbar, text="Issuance slip", style="Accent.TButton", command=lambda: self.open_document_window(document_type="ISSUANCE")).pack(side="left", padx=6)
        ttk.Button(toolbar, text="Add alternate name", command=self.add_alias).pack(side="left", padx=6)
        ttk.Button(toolbar, text="Edit description", command=self.edit_description).pack(side="left", padx=6)
        self.operator_label = ttk.Label(toolbar, text="Operator: Local user", style="ToolbarLabel.TLabel")
        self.operator_label.pack(side="right")

        top = ttk.Frame(self, style="Header.TFrame", padding=(22, 16))
        top.pack(fill="x")
        brand = ttk.Frame(top, style="Header.TFrame"); brand.pack(anchor="w")
        self.logo_image = None
        logo_path = Path(__file__).resolve().parent / "asset" / "warehouse_logo.jpg"
        if Image is not None and logo_path.exists():
            self.logo_image = ImageTk.PhotoImage(Image.open(logo_path).convert("RGB").resize((48, 48)))
            self.iconphoto(True, self.logo_image)
            ttk.Label(brand, image=self.logo_image, background="#172b4d").pack(side="left", padx=(0, 12))
        ttk.Label(brand, text="AIYU Warehouse", style="Header.TLabel").pack(side="left")
        ttk.Label(top, text="Search, review, and capture backlog transactions", style="HeaderSub.TLabel").pack(anchor="w", pady=(2, 10))
        search_bar = ttk.Frame(top, style="Header.TFrame")
        search_bar.pack(fill="x")
        ttk.Label(search_bar, text="Find material", style="HeaderSub.TLabel").pack(side="left")
        self.search_var = tk.StringVar()
        entry = ttk.Entry(search_bar, textvariable=self.search_var, width=30)
        entry.pack(side="left", fill="x", expand=True, padx=10)
        entry.bind("<KeyRelease>", lambda _event: self.after(180, self.search))
        entry.bind("<Return>", lambda _event: self.search())
        ttk.Button(search_bar, text="Search", style="Accent.TButton", command=self.search).pack(side="left")
        self.filter_var = tk.StringVar(value="ALL")
        ttk.Combobox(search_bar, textvariable=self.filter_var, values=["ALL", "REVIEW_REQUIRED", "VALIDATION_LIMITED", "OK"], state="readonly", width=18).pack(side="left", padx=8)
        self.filter_var.trace_add("write", lambda *_: self.search())
        ttk.Button(search_bar, text="Show all", command=self.show_all).pack(side="left", padx=6)
        ttk.Button(search_bar, text="Settings", command=self.open_settings).pack(side="right")
        self.search_entry = entry

        body = ttk.PanedWindow(self, orient="horizontal")
        body.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        left = ttk.Frame(body, style="Panel.TFrame", padding=12)
        right = ttk.Frame(body, style="Panel.TFrame", padding=16)
        body.add(left, weight=1)
        body.add(right, weight=3)

        ttk.Label(left, text="Materials", style="Title.TLabel").pack(anchor="w")
        ttk.Label(left, text="Select a material to inspect its history", style="Muted.TLabel").pack(anchor="w", pady=(2, 10))
        self.listbox = tk.Listbox(left, width=42, exportselection=False, relief="flat", borderwidth=0, bg="white", fg="#1f2937", selectbackground="#dbeafe", selectforeground="#1f2937", activestyle="none", font=("Segoe UI", 10), highlightthickness=1, highlightcolor="#dbe3ee")
        self.listbox.pack(side="left", fill="both", expand=True)
        scroll = ttk.Scrollbar(left, orient="vertical", command=self.listbox.yview)
        scroll.pack(side="right", fill="y")
        self.listbox.configure(yscrollcommand=scroll.set)
        self.listbox.bind("<<ListboxSelect>>", self.select_material)

        self.material_title = ttk.Label(right, text="Select a material", style="Title.TLabel")
        self.material_title.pack(anchor="w")
        self.material_summary = ttk.Label(right, text="", style="Muted.TLabel", wraplength=900)
        self.material_summary.pack(anchor="w", pady=(2, 10))
        table_frame = ttk.Frame(right, style="Panel.TFrame")
        table_frame.pack(fill="both", expand=True)
        columns = ("date", "receive_ref", "received_from", "issue_ref", "issued_to", "qty_in", "qty_out", "balance", "source")
        self.details = ttk.Treeview(table_frame, columns=columns, show="headings")
        headings = {
            "date": "DATE", "receive_ref": "RECEIVING REF.",
            "received_from": "RECEIVED FROM", "issue_ref": "ISSUANCE REF.",
            "issued_to": "ISSUED TO", "qty_in": "IN", "qty_out": "OUT",
            "balance": "BALANCE ON HAND", "source": "SOURCE / STATUS",
        }
        widths = {"date": 92, "receive_ref": 105, "received_from": 145, "issue_ref": 105, "issued_to": 145, "qty_in": 65, "qty_out": 65, "balance": 115, "source": 135}
        for key, label in headings.items():
            self.details.heading(key, text=label, anchor="center")
            self.details.column(key, width=widths[key], anchor="center", stretch=False)
        self.details.tag_configure("receiving", background="#e8f5e9")
        self.details.tag_configure("issuance", background="#fff4cc")
        self.details.tag_configure("opening", background="#e3f2fd")
        self.details.tag_configure("count", background="#e6f4f1")
        self.details.tag_configure("pullout", background="#e0f2fe")
        self.details.tag_configure("paper", background="#fff0d6")
        self.details.tag_configure("issue", foreground="#b45309")
        self.details.pack(side="top", fill="both", expand=True)
        yscroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.details.yview)
        yscroll.pack(side="right", fill="y")
        self.details.configure(yscrollcommand=yscroll.set)
        self.details.bind("<Double-1>", self.edit_selected_entry)
        xscroll = ttk.Scrollbar(table_frame, orient="horizontal", command=self.details.xview)
        xscroll.pack(side="bottom", fill="x")
        self.details.configure(xscrollcommand=xscroll.set)
        warning_bar = ttk.Frame(right, style="Panel.TFrame")
        warning_bar.pack(fill="x", pady=(12, 4))
        ttk.Label(warning_bar, text="Warnings and review notes", style="Muted.TLabel").pack(side="left")
        ttk.Button(warning_bar, text="Mark reviewed", command=self.mark_warnings_reviewed).pack(side="right")
        self.warning_text = tk.Text(right, height=5, wrap="word", state="disabled", relief="flat", bg="#fffaf0", fg="#92400e", padx=10, pady=8, font=("Segoe UI", 9))
        self.warning_text.pack(fill="x")

        form = ttk.Frame(self, style="Panel.TFrame", padding=(10, 4))
        form.pack(fill="x", padx=10, pady=(0, 10))
        ttk.Label(form, text="Need to enter a physical stock-card record?", style="Muted.TLabel").pack(side="left")
        ttk.Button(form, text="Open Add Entry window", style="Accent.TButton", command=self.open_entry_window).pack(side="left", padx=12)
        ttk.Button(form, text="Open Excel source", command=self.open_source).pack(side="left", padx=4)
        ttk.Button(form, text="Export new stock card", style="Accent.TButton", command=self.export_stock_card).pack(side="left", padx=4)
        ttk.Button(form, text="Edit selected entry", command=self.edit_selected_entry).pack(side="left", padx=4)
        ttk.Label(form, text="Editable entries are manual stock-card records.", style="Muted.TLabel").pack(side="left", padx=8)

        self.status = ttk.Label(self, text=f"Database: {self.db_path}", style="Status.TLabel", anchor="w")
        self.status.pack(fill="x", side="bottom")
        self.bind("<Control-f>", lambda _event: (self.search_entry.focus_set(), self.search_entry.select_range(0, "end")))
        self.show_all()

    def _load_settings(self):
        defaults = {
            "source_folder": "",
            "new_stockcards_folder": "",
            "paper_backlog_folder": "",
            "export_folder": "",
            "theme": "Light",
        }
        if self.config_path.exists():
            try:
                saved = json.loads(self.config_path.read_text(encoding="utf-8"))
                defaults.update({key: str(value) for key, value in saved.items() if key in defaults})
            except (OSError, json.JSONDecodeError):
                pass
        return defaults

    def open_settings(self):
        dialog = tk.Toplevel(self)
        dialog.title("AIYU Warehouse Settings")
        dialog.geometry("820x420")
        dialog.minsize(700, 360)
        dialog.transient(self)
        dialog.grab_set()
        dialog.configure(bg="#f3f6fb")
        panel = ttk.Frame(dialog, style="Panel.TFrame", padding=18)
        panel.pack(fill="both", expand=True, padx=12, pady=12)
        ttk.Label(panel, text="Folders and data", style="Title.TLabel").pack(anchor="w")
        ttk.Label(panel, text="Choose the stock-card source and export folders. Source workbooks are read-only.", style="Muted.TLabel").pack(anchor="w", pady=(2, 8))
        variables = {key: tk.StringVar(value=self.settings.get(key, "")) for key in self.settings}
        labels = {
            "source_folder": "Old/source stock cards",
            "new_stockcards_folder": "New working stock cards",
            "paper_backlog_folder": "Paper backlog",
            "export_folder": "Exports",
        }
        for key, label in labels.items():
            row = ttk.Frame(panel, style="Panel.TFrame")
            row.pack(fill="x", pady=2)
            ttk.Label(row, text=label, width=22, style="Body.TLabel").pack(side="left")
            ttk.Entry(row, textvariable=variables[key]).pack(side="left", fill="x", expand=True, padx=8)
            ttk.Button(row, text="Browse", command=lambda var=variables[key]: self._choose_folder(var)).pack(side="right")
            if key in {"source_folder", "export_folder"}:
                ttk.Button(row, text="Open", command=lambda var=variables[key]: self._open_folder(var)).pack(side="right", padx=(0, 5))

        theme_row = ttk.Frame(panel, style="Panel.TFrame"); theme_row.pack(fill="x", pady=2)
        ttk.Label(theme_row, text="Theme", width=22, style="Body.TLabel").pack(side="left")
        ttk.Combobox(theme_row, textvariable=variables["theme"], values=["Light", "Dark"], state="readonly", width=18).pack(side="left", padx=8)

        ttk.Label(panel, text=f"Database: {self.db_path}", style="Muted.TLabel").pack(anchor="w", pady=(6, 0))

        buttons = ttk.Frame(panel, style="Panel.TFrame")
        buttons.pack(fill="x", pady=(8, 0))
        ttk.Button(buttons, text="Cancel", command=dialog.destroy).pack(side="right", padx=(8, 0))
        ttk.Button(buttons, text="Save settings", style="Accent.TButton", command=lambda: self._save_settings(variables, dialog)).pack(side="right")
        ttk.Button(buttons, text="Refresh inventory", command=lambda: (self._save_settings(variables, dialog), self.refresh_inventory())).pack(side="left")
        ttk.Button(buttons, text="Backup database", command=self._backup_database).pack(side="left", padx=6)

    def _choose_folder(self, variable):
        selected = filedialog.askdirectory(parent=self)
        if selected:
            variable.set(selected)

    def _open_folder(self, variable):
        folder = variable.get().strip()
        if not folder or not Path(folder).exists():
            messagebox.showwarning("Folder unavailable", "Choose an existing folder first.", parent=self)
            return
        try:
            os.startfile(folder)
        except OSError as exc:
            messagebox.showerror("Could not open folder", str(exc), parent=self)

    def _save_settings(self, variables, dialog):
        self.settings = {key: variable.get().strip() for key, variable in variables.items()}
        self.config_path.write_text(json.dumps(self.settings, indent=2), encoding="utf-8")
        self._style()
        dialog.destroy()
        self.status.configure(text=f"Settings saved  •  Source folder: {self.settings['source_folder'] or '[not configured]'}")

    def _backup_database(self):
        target = filedialog.asksaveasfilename(parent=self, title="Backup AIYU Warehouse database", defaultextension=".sqlite3", filetypes=[("SQLite database", "*.sqlite3"), ("All files", "*.*")])
        if not target: return
        try:
            shutil.copy2(self.db_path, target); messagebox.showinfo("Backup created", f"Database backup saved to:\n{target}", parent=self)
        except OSError as exc: messagebox.showerror("Backup failed", str(exc), parent=self)

    def _style(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        dark = self.settings.get("theme", "Light") == "Dark"
        panel, fg, muted, toolbar = ("#1f2937", "#f8fafc", "#cbd5e1", "#111827") if dark else ("white", "#1f2937", "#64748b", "#e8eef7")
        style.configure("Header.TFrame", background="#172b4d")
        style.configure("Header.TLabel", background="#172b4d", foreground="white", font=("Segoe UI", 19, "bold"))
        style.configure("HeaderSub.TLabel", background="#172b4d", foreground="#b9c9df", font=("Segoe UI", 9))
        style.configure("Toolbar.TFrame", background=toolbar)
        style.configure("ToolbarLabel.TLabel", background=toolbar, foreground=muted, font=("Segoe UI", 9))
        style.configure("Panel.TFrame", background=panel)
        style.configure("Title.TLabel", background=panel, foreground=fg, font=("Segoe UI", 13, "bold"))
        style.configure("Body.TLabel", background=panel, foreground=fg, font=("Segoe UI", 9))
        style.configure("Muted.TLabel", background=panel, foreground=muted, font=("Segoe UI", 9))
        style.configure("Status.TLabel", background="#e8eef7", foreground="#53657d", padding=(16, 7), font=("Segoe UI", 9))
        style.configure("Panel.TLabelframe", background="white", foreground="#1f2937", bordercolor="#dbe3ee")
        style.configure("Panel.TLabelframe.Label", background="white", foreground="#1f2937", font=("Segoe UI", 10, "bold"))
        style.configure("Accent.TButton", background="#2563eb", foreground="white", padding=(12, 7), font=("Segoe UI", 9, "bold"))
        style.map("Accent.TButton", background=[("active", "#1d4ed8")])
        style.configure("TButton", padding=(9, 6), font=("Segoe UI", 9))
        style.configure("TLabel", font=("Segoe UI", 9))
        style.configure("TEntry", padding=6)
        style.map("TEntry", fieldbackground=[("disabled", "#e5e7eb")], foreground=[("disabled", "#9ca3af")])
        style.map("TCombobox", fieldbackground=[("disabled", "#e5e7eb")], foreground=[("disabled", "#9ca3af")])
        style.configure("Treeview", background="white", fieldbackground="white", foreground="#1f2937", rowheight=30, font=("Segoe UI", 9), borderwidth=1, relief="solid")
        style.configure("Treeview.Heading", background="#e8eef7", foreground="#334155", font=("Segoe UI", 9, "bold"), padding=8, relief="raised")

    def attach_material_autocomplete(self, entry, variable, parent):
        popup = {"win": None}
        def update(_event=None):
            term = variable.get().strip().upper()
            if not term:
                if popup["win"]: popup["win"].destroy(); popup["win"] = None
                return
            conn = connect(self.db_path)
            rows = conn.execute("SELECT m.name, COALESCE(m.description,'') AS description, COALESCE((SELECT group_concat(alias,' ') FROM material_aliases a WHERE a.material_id=m.id),'') AS aliases FROM materials m WHERE UPPER(m.name) LIKE ? OR UPPER(COALESCE(m.description,'')) LIKE ? OR EXISTS (SELECT 1 FROM material_aliases a WHERE a.material_id=m.id AND UPPER(a.alias) LIKE ?) ORDER BY m.name LIMIT 8", (f"%{term}%", f"%{term}%", f"%{term}%")).fetchall(); conn.close()
            if popup["win"]: popup["win"].destroy()
            if not rows: popup["win"] = None; return
            win = tk.Toplevel(parent); popup["win"] = win; win.overrideredirect(True); win.attributes("-topmost", True)
            box = tk.Listbox(win, width=55, height=min(8,len(rows)), exportselection=False)
            for r in rows: box.insert("end", f"{r['name']}  ·  {r['description'] or r['aliases']}")
            box.pack()
            x,y=entry.winfo_rootx(),entry.winfo_rooty()+entry.winfo_height(); win.geometry(f"+{x}+{y}")
            def choose(_=None):
                if box.curselection(): variable.set(rows[box.curselection()[0]]["name"])
                win.destroy(); popup["win"] = None
            box.bind("<ButtonRelease-1>", choose); box.bind("<Return>", choose)
        entry.bind("<KeyRelease>", update)

    def attach_project_autocomplete(self, entry, variable, parent):
        popup = {"win": None}
        def update(_event=None):
            term = variable.get().strip().upper()
            if not term:
                if popup["win"]: popup["win"].destroy(); popup["win"] = None
                return
            conn = connect(self.db_path)
            rows = conn.execute("SELECT name FROM projects WHERE normalized_name LIKE ? ORDER BY name LIMIT 8", (f"%{term}%",)).fetchall(); conn.close()
            if popup["win"]: popup["win"].destroy()
            if not rows: popup["win"] = None; return
            win = tk.Toplevel(parent); popup["win"] = win; win.overrideredirect(True); win.attributes("-topmost", True)
            box = tk.Listbox(win, width=55, height=min(8, len(rows)), exportselection=False)
            for row in rows: box.insert("end", row["name"])
            box.pack(); win.geometry(f"+{entry.winfo_rootx()}+{entry.winfo_rooty()+entry.winfo_height()}")
            def choose(_=None):
                if box.curselection(): variable.set(rows[box.curselection()[0]]["name"])
                win.destroy(); popup["win"] = None
            box.bind("<ButtonRelease-1>", choose); box.bind("<Return>", choose)
        entry.bind("<KeyRelease>", update)

    def attach_supplier_autocomplete(self, entry, variable, parent):
        popup = {"win": None}
        def update(_event=None):
            term = variable.get().strip().upper()
            if not term:
                if popup["win"]: popup["win"].destroy(); popup["win"] = None
                return
            conn = connect(self.db_path)
            rows = conn.execute("SELECT supplier, COUNT(*) AS uses FROM transactions WHERE TRIM(COALESCE(supplier,'')) <> '' AND UPPER(supplier) LIKE ? GROUP BY supplier ORDER BY uses DESC, supplier LIMIT 8", (f"%{term}%",)).fetchall(); conn.close()
            if popup["win"]: popup["win"].destroy()
            if not rows: popup["win"] = None; return
            win = tk.Toplevel(parent); popup["win"] = win; win.overrideredirect(True); win.attributes("-topmost", True)
            box = tk.Listbox(win, width=45, height=min(8, len(rows)), exportselection=False)
            for row in rows: box.insert("end", row["supplier"])
            box.pack(); win.geometry(f"+{entry.winfo_rootx()}+{entry.winfo_rooty()+entry.winfo_height()}")
            def choose(_=None):
                if box.curselection(): variable.set(rows[box.curselection()[0]]["supplier"])
                win.destroy(); popup["win"] = None
            box.bind("<ButtonRelease-1>", choose); box.bind("<Return>", choose)
        entry.bind("<KeyRelease>", update)

    def attach_recurring_autocomplete(self, entry, variable, parent, columns):
        popup = {"win": None}
        allowed = {"supplier", "driver", "vehicle", "received_by", "released_by", "document_number", "po_number", "reference_number"}
        columns = [column for column in columns if column in allowed]
        def update(_event=None):
            term = variable.get().strip().upper()
            if not term or not columns:
                if popup["win"]: popup["win"].destroy(); popup["win"] = None
                return
            conn = connect(self.db_path); clauses=[]; params=[]
            for column in columns:
                table = "documents" if column in {"driver", "vehicle", "received_by", "released_by", "document_number", "po_number"} else "transactions"
                clauses.append(f"SELECT {column} AS value FROM {table} WHERE TRIM(COALESCE({column},'')) <> '' AND UPPER({column}) LIKE ?")
                params.append(f"%{term}%")
            rows = conn.execute("SELECT value, COUNT(*) AS uses FROM (" + " UNION ALL ".join(clauses) + ") GROUP BY value ORDER BY uses DESC, value LIMIT 8", params).fetchall(); conn.close()
            if popup["win"]: popup["win"].destroy()
            if not rows: popup["win"] = None; return
            win = tk.Toplevel(parent); popup["win"] = win; win.overrideredirect(True); win.attributes("-topmost", True)
            box = tk.Listbox(win, width=42, height=min(8, len(rows)), exportselection=False)
            for row in rows: box.insert("end", row["value"])
            box.pack(); win.geometry(f"+{entry.winfo_rootx()}+{entry.winfo_rooty()+entry.winfo_height()}")
            def choose(_=None):
                if box.curselection(): variable.set(rows[box.curselection()[0]]["value"])
                win.destroy(); popup["win"] = None
            box.bind("<ButtonRelease-1>", choose); box.bind("<Return>", choose)
        entry.bind("<KeyRelease>", update)

    def search(self):
        term = self.search_var.get().strip().upper()
        conn = connect(self.db_path)
        try:
            if term:
                rows = conn.execute("SELECT m.id, m.name, m.normalized_name, m.location, m.description, m.transaction_count, m.last_status, COALESCE((SELECT group_concat(alias, ' ') FROM material_aliases a WHERE a.material_id=m.id), '') AS aliases, COALESCE((SELECT group_concat(reference_number, ' ') FROM transactions t WHERE t.material_id=m.id), '') AS transaction_refs, COALESCE((SELECT group_concat(d.po_number || ' ' || d.document_number, ' ') FROM documents d JOIN document_lines dl ON dl.document_id=d.id WHERE dl.material_id=m.id), '') AS document_refs FROM materials m ORDER BY m.name").fetchall()
            else:
                rows = conn.execute("SELECT m.id, m.name, m.normalized_name, m.location, m.description, m.transaction_count, m.last_status, COALESCE((SELECT group_concat(alias, ' ') FROM material_aliases a WHERE a.material_id=m.id), '') AS aliases, '' AS transaction_refs, '' AS document_refs FROM materials m ORDER BY m.transaction_count DESC, m.name LIMIT 100").fetchall()
        except sqlite3.OperationalError:
            rows = conn.execute("SELECT m.id, m.name, m.normalized_name, m.location, m.description, m.transaction_count, m.last_status, COALESCE((SELECT group_concat(alias, ' ') FROM material_aliases a WHERE a.material_id=m.id), '') AS aliases, '' AS transaction_refs, '' AS document_refs FROM materials m").fetchall()
        conn.close()
        tokens = [token for token in term.split() if token]
        matches = []
        for row in rows:
            name = (row["normalized_name"] or "").upper()
            location = (row["location"] or "").upper()
            searchable = f"{name} {location} {(row['aliases'] or '').upper()} {(row['description'] or '').upper()} {(row['transaction_refs'] or '').upper()} {(row['document_refs'] or '').upper()}"
            if tokens and not all(token in searchable for token in tokens):
                continue
            if not tokens:
                score = 0
            elif name == term:
                score = 0
            elif name.startswith(term):
                score = 1
            elif all(name.startswith(token) for token in tokens):
                score = 2
            elif term in name:
                score = 3
            else:
                score = 4
            matches.append((score, name, row))
        matches.sort(key=lambda item: (item[0], item[1]))
        rows = [item[2] for item in matches]
        selected_filter = self.filter_var.get()
        if selected_filter != "ALL":
            rows = [row for row in rows if row["last_status"] == selected_filter]
        rows = rows[:100]
        self.listbox.delete(0, "end")
        self.material_ids = []
        for row in rows:
            self.material_ids.append(int(row["id"]))
            self.listbox.insert("end", row["name"])
        self.status.configure(text=f"{len(rows)} material matches")

    def show_all(self):
        self.search_var.set("")
        self.search()

    def refresh_inventory(self):
        source = self.settings.get("source_folder", "").strip()
        if not source:
            messagebox.showwarning("Source folder required", "Select the stock-card source folder in Settings first.")
            return
        source_signature = self._source_signature(Path(source))
        conn = connect(self.db_path)
        cached = conn.execute("SELECT value FROM metadata WHERE key='source_signature'").fetchone()
        conn.close()
        if cached and cached["value"] == source_signature:
            self.search(); self.status.configure(text="No source changes detected • Inventory is current")
            return
        runtime_root = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
        if getattr(sys, "frozen", False):
            batch_script = Path(sys.executable).resolve().parent / "AIYU Scanner" / "AIYU Scanner.exe"
            backend_script = Path(sys.executable).resolve().parent / "AIYU Database Builder" / "AIYU Database Builder.exe"
            scan_command = [str(batch_script), source, "--output", str(DEFAULT_BATCH)]
            build_command = [str(backend_script), "--batch", str(DEFAULT_BATCH), "--db", str(self.db_path), "build"]
        else:
            batch_script = runtime_root / "stock_card_batch_v1.py"
            backend_script = runtime_root / "stock_card_backend_v1.py"
            scan_command = [sys.executable, str(batch_script), source, "--output", str(DEFAULT_BATCH)]
            build_command = [sys.executable, str(backend_script), "--batch", str(DEFAULT_BATCH), "--db", str(self.db_path), "build"]
        try:
            self.status.configure(text="Reading stock cards..."); self.update_idletasks()
            subprocess.run(scan_command, check=True)
            self.status.configure(text="Building inventory database..."); self.update_idletasks()
            subprocess.run(build_command, check=True)
            conn = connect(self.db_path)
            conn.execute("INSERT INTO metadata(key,value) VALUES ('source_signature',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (source_signature,))
            conn.commit(); conn.close()
            self.search(); self.status.configure(text="Inventory refreshed successfully")
        except (OSError, subprocess.CalledProcessError) as exc:
            messagebox.showerror("Refresh failed", str(exc))

    def _source_signature(self, source: Path) -> str:
        digest = hashlib.sha256()
        if source.is_file():
            paths = [source]
        elif source.is_dir():
            excluded = DEFAULT_BATCH.resolve()
            paths = []
            for path in source.rglob("*"):
                if not path.is_file() or path.name.startswith("~$") or path.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}: continue
                try: path.resolve().relative_to(excluded); continue
                except ValueError: paths.append(path)
            paths.sort()
        else:
            return "missing"
        for path in paths:
            stat = path.stat()
            digest.update(str(path.resolve()).encode("utf-8")); digest.update(str(stat.st_size).encode("ascii")); digest.update(str(stat.st_mtime_ns).encode("ascii"))
        return digest.hexdigest()

    def select_material(self, _event=None):
        selection = self.listbox.curselection()
        if not selection:
            return
        self.selected_id = self.material_ids[selection[0]]
        conn = connect(self.db_path)
        material = conn.execute("SELECT * FROM materials WHERE id=?", (self.selected_id,)).fetchone()
        aliases = conn.execute("SELECT alias FROM material_aliases WHERE material_id=? ORDER BY alias", (self.selected_id,)).fetchall()
        sources = conn.execute("SELECT * FROM sources WHERE material_id=? ORDER BY id", (self.selected_id,)).fetchall()
        transactions = conn.execute("SELECT * FROM transactions WHERE material_id=? ORDER BY date, id", (self.selected_id,)).fetchall()
        open_issues = conn.execute("SELECT issue_code, source_id FROM issues WHERE material_id=? AND status='OPEN' ORDER BY id", (self.selected_id,)).fetchall()
        conn.close()
        self.material_title.configure(text=material["name"])
        alias_text = ", ".join(row["alias"] for row in aliases)
        summary = material["description"] or "No description yet."
        if material["location"]:
            summary += f"  •  Location: {material['location']}"
        if alias_text:
            summary += f"  •  Also known as: {alias_text}"
        self.material_summary.configure(text=summary)
        self.details.delete(*self.details.get_children())
        for row_index, row in enumerate(transactions):
            tx_type = (row["transaction_type"] or "").upper()
            receive_ref = row["reference_number"] if tx_type == "RECEIVING" else ""
            issue_ref = row["reference_number"] if tx_type == "ISSUANCE" else ""
            qty_in = row["quantity_in"] if row["quantity_in"] is not None else ""
            qty_out = row["quantity_out"] if row["quantity_out"] is not None else ""
            if row["origin"] == "PAPER":
                tags = ("paper",)
            elif tx_type == "RECEIVING":
                tags = ("receiving",)
            elif tx_type == "ISSUANCE":
                tags = ("issuance",)
            elif tx_type == "OPENING_BALANCE":
                tags = ("opening",)
            elif tx_type == "PHYSICAL_COUNT":
                tags = ("count",)
            elif tx_type == "PULL_OUT":
                tags = ("pullout",)
            else:
                tags = ()
            self.details.insert("", "end", iid=f"tx_{row['id']}", values=(
                row["date"], receive_ref, row["supplier"] or "", issue_ref,
                row["project"] or "", qty_in, qty_out,
                row["recorded_balance"] if row["recorded_balance"] is not None else "",
                f"{row['origin']} / {row['status']}",
            ), tags=tags)
        source_by_id = {source["id"]: source for source in sources}
        warnings = []
        for issue in open_issues:
            source = source_by_id.get(issue["source_id"])
            context = f"{source['source_file']} / block {source['block_number']} side {source['side_number']}" if source else "material record"
            warnings.append(f"{issue['issue_code']}  ·  {context}")
        self.warning_text.configure(state="normal")
        self.warning_text.delete("1.0", "end")
        self.warning_text.insert("1.0", "\n".join(warnings) if warnings else "No warnings recorded for this material.")
        self.warning_text.configure(state="disabled")
        self.status.configure(text=f"{material['name']}  •  {len(sources)} source card(s)  •  {len(transactions)} transaction(s)")
        if sources:
            source = sources[0]
            self.selected_source = (
                source["source_file"], source["source_sheet"],
                source["block_number"], source["side_number"],
                0,
            )
        else:
            self.selected_source = None

    def add_alias(self):
        if self.selected_id is None:
            messagebox.showwarning("Select a material", "Select the main material first.")
            return
        dialog = tk.Toplevel(self); dialog.title("Add alternate material name"); dialog.geometry("430x180"); dialog.transient(self); dialog.grab_set()
        panel = ttk.Frame(dialog, style="Panel.TFrame", padding=18); panel.pack(fill="both", expand=True)
        ttk.Label(panel, text="Alternate name or acronym", style="Title.TLabel").pack(anchor="w")
        ttk.Label(panel, text="Example: CWN for Common Wire Nails. This keeps the original record and links the search name safely.", style="Muted.TLabel", wraplength=380).pack(anchor="w", pady=(4, 10))
        value = tk.StringVar(); field = ttk.Entry(panel, textvariable=value); field.pack(fill="x"); field.focus_set()
        def save():
            try: add_material_alias(self.db_path, self.selected_id, value.get())
            except Exception as exc: messagebox.showerror("Could not add alternate name", str(exc), parent=dialog); return
            dialog.destroy(); self.search(); self.status.configure(text="Saved successfully • Alternate name added")
        ttk.Button(panel, text="Save alternate name", style="Accent.TButton", command=save).pack(anchor="e", pady=(14, 0))

    def edit_description(self):
        if self.selected_id is None:
            messagebox.showwarning("Select a material", "Select a material first.")
            return
        conn = connect(self.db_path); row = conn.execute("SELECT name, description FROM materials WHERE id=?", (self.selected_id,)).fetchone(); conn.close()
        dialog = tk.Toplevel(self); dialog.title("Edit material description"); dialog.geometry("500x220"); dialog.transient(self); dialog.grab_set()
        panel = ttk.Frame(dialog, style="Panel.TFrame", padding=18); panel.pack(fill="both", expand=True)
        ttk.Label(panel, text=f"Description for {row['name']}", style="Title.TLabel").pack(anchor="w")
        value = tk.StringVar(value=row["description"] or ""); ttk.Entry(panel, textvariable=value).pack(fill="x", pady=(12, 0));
        def save():
            set_material_description(self.db_path, self.selected_id, value.get()); dialog.destroy(); self.select_material()
        ttk.Button(panel, text="Save description", style="Accent.TButton", command=save).pack(anchor="e", pady=(14, 0))

    def open_source(self):
        if not self.selected_source:
            messagebox.showwarning("No source selected", "Select a material with an Excel source first.")
            return
        source_file, sheet_name, block_number, side_number, _ = self.selected_source
        conn = connect(self.db_path)
        row = conn.execute("SELECT value FROM metadata WHERE key='source_root'").fetchone()
        conn.close()
        configured_root = self.settings.get("source_folder", "").strip()
        source_root = configured_root or (row["value"] if row else "")
        if not source_root:
            messagebox.showerror("Source folder unavailable", "This database does not contain the original source folder path.")
            return
        workbook = Path(source_root) / source_file
        if not workbook.exists():
            messagebox.showerror("Workbook not found", f"Expected source workbook was not found:\n{workbook}\n\nSheet: {sheet_name}\nBlock: {block_number}, side {side_number}")
            return
        try:
            os.startfile(str(workbook))
            self.status.configure(text=f"Opened {workbook.name}  •  sheet {sheet_name}  •  block {block_number}, side {side_number}")
        except OSError as exc:
            messagebox.showerror("Could not open workbook", str(exc))

    def export_stock_card(self):
        if self.selected_id is None:
            messagebox.showwarning("Select a material", "Select a material before exporting a stock card.")
            return
        conn = connect(self.db_path)
        material = conn.execute("SELECT * FROM materials WHERE id=?", (self.selected_id,)).fetchone()
        transactions = conn.execute("SELECT * FROM transactions WHERE material_id=? ORDER BY date, id", (self.selected_id,)).fetchall()
        conn.close()
        if not material:
            messagebox.showerror("Material unavailable", "The selected material could not be found.")
            return
        export_root = Path(self.settings.get("export_folder") or (Path(__file__).resolve().parent / "05_Exports"))
        export_root.mkdir(parents=True, exist_ok=True)
        safe_name = re.sub(r"[^A-Za-z0-9._ -]+", "", material["name"]).strip() or "stock_card"
        output_path = export_root / f"{safe_name}_stock_card.xlsx"
        template_path = Path(__file__).resolve().parent / "TemplateForms" / "STOCK CARD.xlsx"
        wb = load_workbook(template_path) if template_path.exists() else Workbook()
        ws = wb.active
        ws.title = "ITEM"
        thin = Side(style="thin", color="B7C3D0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        blue = PatternFill("solid", fgColor="D9EAF7")
        ws["B4"] = material["name"]
        ws["B5"] = material["location"] or ""
        # The approved template already contains the merged, formatted headers.
        balance = None
        for row_index, tx in enumerate(transactions, 9):
            tx_type = (tx["transaction_type"] or "").upper()
            qty_in = tx["quantity_in"]
            qty_out = tx["quantity_out"]
            if tx["net_movement"] is not None and balance is not None:
                balance += tx["net_movement"]
            elif tx["recorded_balance"] is not None:
                balance = tx["recorded_balance"]
            values = [
                tx["date"],
                tx["reference_number"] if tx_type == "RECEIVING" else "",
                tx["supplier"] or "",
                tx["reference_number"] if tx_type == "ISSUANCE" else "",
                tx["project"] or "",
                qty_in or "",
                qty_out or "",
                tx["recorded_balance"] if tx["recorded_balance"] is not None else (balance if balance is not None else ""),
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row_index, col, value)
                # Keep the approved template's existing border, alignment, and fills.
        wb.save(output_path)
        note = ""
        messagebox.showinfo("Stock card exported", f"Created:\n{output_path}{note}\n\nThe original Excel source was not changed.")

    def export_all_stock_cards(self):
        target = filedialog.askdirectory(parent=self, title="Choose folder for all stock-card exports")
        if not target: return
        conn = connect(self.db_path)
        materials = conn.execute("SELECT id FROM materials ORDER BY name").fetchall()
        conn.close()
        if not materials:
            messagebox.showwarning("No materials", "There are no materials to export.", parent=self); return
        progress = tk.Toplevel(self); progress.title("Exporting stock cards"); progress.geometry("520x150"); progress.transient(self); progress.grab_set()
        progress.protocol("WM_DELETE_WINDOW", lambda: None)
        progress_panel = ttk.Frame(progress, padding=18); progress_panel.pack(fill="both", expand=True)
        progress_label = ttk.Label(progress_panel, text="Preparing stock-card exports..."); progress_label.pack(anchor="w", pady=(0, 10))
        progress_bar = ttk.Progressbar(progress_panel, mode="determinate", maximum=len(materials)); progress_bar.pack(fill="x")
        progress_count = ttk.Label(progress_panel, text=f"0 of {len(materials)} completed"); progress_count.pack(anchor="w", pady=(8, 0))
        progress.update_idletasks()
        previous = self.selected_id; exported = 0
        try:
            for row in materials:
                self.selected_id = int(row["id"])
                self.settings["export_folder"] = target
                name_conn = connect(self.db_path); material_name = name_conn.execute("SELECT name FROM materials WHERE id=?", (self.selected_id,)).fetchone()["name"]; name_conn.close()
                progress_label.configure(text=f"Exporting: {material_name}")
                self._export_selected_stock_card_silent()
                exported += 1
                progress_bar.configure(value=exported); progress_count.configure(text=f"{exported} of {len(materials)} completed"); progress.update_idletasks()
        except Exception as exc:
            progress.destroy()
            messagebox.showerror("Export all failed", str(exc), parent=self); return
        finally:
            self.selected_id = previous
        progress.destroy()
        self.status.configure(text=f"Exported successfully • {exported} stock cards")

    def _export_selected_stock_card_silent(self):
        conn = connect(self.db_path)
        material = conn.execute("SELECT * FROM materials WHERE id=?", (self.selected_id,)).fetchone()
        transactions = conn.execute("SELECT * FROM transactions WHERE material_id=? ORDER BY date, id", (self.selected_id,)).fetchall()
        conn.close()
        if not material: return
        export_root = Path(self.settings.get("export_folder") or (Path(__file__).resolve().parent / "05_Exports")); export_root.mkdir(parents=True, exist_ok=True)
        safe_name = re.sub(r"[^A-Za-z0-9._ -]+", "", material["name"]).strip() or "stock_card"
        template_path = Path(__file__).resolve().parent / "TemplateForms" / "STOCK CARD.xlsx"
        wb = load_workbook(template_path) if template_path.exists() else Workbook(); ws = wb.active; ws.title = "ITEM"
        ws["B4"] = material["name"]; ws["B5"] = material["location"] or ""; balance = None
        for row_index, tx in enumerate(transactions, 9):
            tx_type = (tx["transaction_type"] or "").upper(); qty_in = tx["quantity_in"]; qty_out = tx["quantity_out"]
            if tx["net_movement"] is not None and balance is not None: balance += tx["net_movement"]
            elif tx["recorded_balance"] is not None: balance = tx["recorded_balance"]
            for col, value in enumerate([tx["date"], tx["reference_number"] if tx_type == "RECEIVING" else "", tx["supplier"] or "", tx["reference_number"] if tx_type == "ISSUANCE" else "", tx["project"] or "", qty_in or "", qty_out or "", tx["recorded_balance"] if tx["recorded_balance"] is not None else (balance if balance is not None else "")], 1): ws.cell(row_index, col, value)
        wb.save(export_root / f"{safe_name}_stock_card.xlsx")

    def mark_warnings_reviewed(self):
        if self.selected_id is None:
            messagebox.showwarning("Select a material", "Select a material first.")
            return
        conn = connect(self.db_path)
        conn.execute("UPDATE issues SET status='REVIEWED' WHERE material_id=? AND status='OPEN'", (self.selected_id,))
        conn.commit(); conn.close()
        self.select_material()
        messagebox.showinfo("Warnings updated", "Open warnings for this material are now marked REVIEWED.")

    def edit_selected_entry(self, _event=None):
        selected = self.details.selection()
        if not selected or not selected[0].startswith("tx_"):
            messagebox.showwarning("Select an entry", "Select a stock-card row first.", parent=self); return
        tx_id = int(selected[0][3:]); conn = connect(self.db_path); row = conn.execute("SELECT * FROM transactions WHERE id=? AND origin='PAPER'", (tx_id,)).fetchone(); conn.close()
        if not row:
            messagebox.showinfo("Imported stock-card record", "Only manually entered backlog records can be edited here. Imported Excel records remain read-only.", parent=self); return
        dialog = tk.Toplevel(self); dialog.title("Edit stock card entry"); dialog.transient(self); dialog.grab_set()
        form = ttk.Frame(dialog, padding=16); form.pack(fill="both", expand=True)
        values = {key: tk.StringVar(value=str(row[key] or "")) for key in ("date", "reference_number", "quantity", "unit", "supplier", "project", "notes")}
        values["transaction_type"] = tk.StringVar(value=row["transaction_type"] or "RECEIVING")
        fields = [("Date", "date"), ("Type", "transaction_type"), ("PO / RR / IS number", "reference_number"), ("Quantity", "quantity"), ("Unit", "unit"), ("Received from", "supplier"), ("Project/site", "project"), ("Issues / remarks", "notes")]
        for i, (label, key) in enumerate(fields):
            ttk.Label(form, text=label).grid(row=i, column=0, sticky="w", padx=(0, 8), pady=4)
            if key == "date": widget = DateEntry(form, textvariable=values[key], date_pattern="yyyy-mm-dd", width=25)
            elif key == "transaction_type": widget = ttk.Combobox(form, textvariable=values[key], values=["RECEIVING", "ISSUANCE", "PULL_OUT", "OPENING_BALANCE", "PHYSICAL_COUNT"], state="readonly", width=23)
            else: widget = ttk.Entry(form, textvariable=values[key], width=26)
            widget.grid(row=i, column=1, sticky="ew", pady=4)
            if key == "project": self.attach_project_autocomplete(widget, values[key], dialog)
            if key == "supplier": self.attach_supplier_autocomplete(widget, values[key], dialog)
        def save_edit():
            try:
                register_project(self.db_path, values["project"].get())
                update_paper_record(self.db_path, tx_id, values["date"].get(), values["transaction_type"].get(), float(values["quantity"].get()), values["reference_number"].get() or None, values["unit"].get() or None, values["supplier"].get() or None, values["project"].get() or None, values["notes"].get() or None)
            except Exception as exc: messagebox.showerror("Could not update entry", str(exc), parent=dialog); return
            dialog.destroy(); self.show_all(); self.status.configure(text="Saved successfully • Stock-card entry updated")
        ttk.Button(form, text="Save changes", style="Accent.TButton", command=save_edit).grid(row=len(fields), column=0, columnspan=2, sticky="ew", pady=(10, 0))

    def open_entry_window(self, preset_type="RECEIVING"):
        dialog = tk.Toplevel(self)
        dialog.title("Add stock-card entry")
        dialog.geometry("980x760")
        dialog.minsize(850, 650)
        dialog.transient(self)
        dialog.grab_set()
        dialog.configure(bg="#f3f6fb")

        top = ttk.Frame(dialog, style="Header.TFrame", padding=(18, 12)); top.pack(fill="x")
        ttk.Label(top, text="Add stock-card entry", style="Header.TLabel").pack(anchor="w")
        ttk.Label(top, text="Find an existing material or create a new one", style="HeaderSub.TLabel").pack(anchor="w")

        body = ttk.Frame(dialog, style="Panel.TFrame", padding=14); body.pack(fill="both", expand=True, padx=12, pady=12)
        left = ttk.Frame(body, style="Panel.TFrame"); left.pack(side="left", fill="both", expand=True, padx=(0, 14))
        right = ttk.Frame(body, style="Panel.TFrame", padding=12); right.pack(side="right", fill="y")
        ttk.Label(left, text="Existing materials", style="Title.TLabel").pack(anchor="w")
        find_var = tk.StringVar(); find_entry = ttk.Entry(left, textvariable=find_var); find_entry.pack(fill="x", pady=(8, 8))
        material_tree = ttk.Treeview(left, columns=("name", "location"), show="headings", selectmode="browse")
        material_tree.heading("name", text="ITEM"); material_tree.heading("location", text="LOCATION")
        material_tree.column("name", width=300); material_tree.column("location", width=180)
        material_tree.pack(fill="both", expand=True)
        selected = {"id": None}

        def load_choices(*_):
            conn = connect(self.db_path); term = find_var.get().strip().upper()
            # Search the complete material table; ranking below limits only visible results.
            rows = conn.execute("SELECT m.id,m.name,m.normalized_name,m.location,m.description,COALESCE((SELECT group_concat(alias, ' ') FROM material_aliases a WHERE a.material_id=m.id), '') AS aliases FROM materials m ORDER BY m.name").fetchall()
            tokens = [token for token in term.split() if token]
            ranked = []
            for row in rows:
                name = (row["normalized_name"] or "").upper()
                location = (row["location"] or "").upper()
                searchable = f"{name} {location} {(row['aliases'] or '').upper()} {(row['description'] or '').upper()}"
                if tokens and not all(token in searchable for token in tokens): continue
                score = 0 if not tokens or name == term else 1 if name.startswith(term) else 2 if term in name else 3
                ranked.append((score, name, row))
            ranked.sort(key=lambda item: (item[0], item[1]))
            material_tree.delete(*material_tree.get_children())
            for _, _, row in ranked[:100]: material_tree.insert("", "end", iid=str(row["id"]), values=(row["name"], row["location"] or "[blank]"))
            conn.close()
            return
            try:
                if term:
                    fts_term = " ".join(f'"{token.replace(chr(34), "")}"*' for token in term.split() if token)
                    rows = conn.execute("SELECT m.id,m.name,m.normalized_name,m.location,m.description,COALESCE(f.aliases,'') AS aliases FROM materials m LEFT JOIN materials_fts f ON m.id=CAST(f.material_id AS INTEGER) WHERE (f.material_id IS NOT NULL AND materials_fts MATCH ?) OR UPPER(COALESCE(m.description,'')) LIKE ? ORDER BY m.name LIMIT 100", (fts_term, f"%{term}%")).fetchall()
                else:
                    rows = conn.execute("SELECT m.id,m.name,m.normalized_name,m.location,m.description,COALESCE((SELECT group_concat(alias, ' ') FROM material_aliases a WHERE a.material_id=m.id), '') AS aliases FROM materials m ORDER BY m.name LIMIT 100").fetchall()
            except sqlite3.OperationalError:
                rows = conn.execute("SELECT m.id,m.name,m.normalized_name,m.location,m.description,COALESCE((SELECT group_concat(alias, ' ') FROM material_aliases a WHERE a.material_id=m.id), '') AS aliases FROM materials m ORDER BY m.name LIMIT 100").fetchall()
            conn.close()
            tokens = [token for token in term.split() if token]
            ranked = []
            for row in rows:
                name = (row["normalized_name"] or "").upper()
                location = (row["location"] or "").upper()
                searchable = f"{name} {location} {(row['aliases'] or '').upper()} {(row['description'] or '').upper()}"
                if tokens and not all(token in searchable for token in tokens):
                    continue
                score = 0 if not tokens or name == term else 1 if name.startswith(term) else 2 if term in name else 3
                ranked.append((score, name, row))
            ranked.sort(key=lambda item: (item[0], item[1]))
            rows = [item[2] for item in ranked]
            rows = rows[:100]
            material_tree.delete(*material_tree.get_children())
            for row in rows: material_tree.insert("", "end", iid=str(row["id"]), values=(row["name"], row["location"] or "[blank]"))
        find_var.trace_add("write", load_choices); load_choices()

        ttk.Label(right, text="Transaction", style="Title.TLabel").grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 12))
        name_var, location_var, description_var = tk.StringVar(), tk.StringVar(), tk.StringVar()
        date_var, type_var, ref_var = tk.StringVar(), tk.StringVar(value=preset_type), tk.StringVar()
        qty_var, unit_var, supplier_var, project_var, notes_var = (tk.StringVar() for _ in range(5))
        fields = [("New item name", name_var), ("Short description", description_var), ("Location", location_var), ("Date", date_var), ("Type", type_var), ("PO / RR / IS number", ref_var), ("Quantity", qty_var), ("Unit", unit_var), ("Received from", supplier_var), ("Project/site", project_var), ("Issues / remarks", notes_var)]
        input_widgets = {}
        for row, (label, var) in enumerate(fields, 1):
            ttk.Label(right, text=label, style="Muted.TLabel").grid(row=row, column=0, sticky="w", padx=(0, 8), pady=4)
            if label == "Type":
                widget = ttk.Combobox(right, textvariable=var, values=["RECEIVING", "ISSUANCE", "PULL_OUT", "OPENING_BALANCE", "PHYSICAL_COUNT"], state="readonly", width=25)
            elif label == "Date":
                widget = DateEntry(right, textvariable=var, width=25, date_pattern="yyyy-mm-dd")
            else:
                widget = ttk.Entry(right, textvariable=var, width=28)
            widget.grid(row=row, column=1, sticky="ew", pady=4)
            input_widgets[label] = widget
            if label == "Project/site": self.attach_project_autocomplete(widget, var, dialog)
            if label == "Received from": self.attach_supplier_autocomplete(widget, var, dialog)
            if label == "PO / RR / IS number": self.attach_recurring_autocomplete(widget, var, dialog, ["reference_number", "document_number", "po_number"])

        def update_type_fields(*_):
            receiving = type_var.get() in {"RECEIVING", "PULL_OUT"}
            issuance = type_var.get() == "ISSUANCE"
            input_widgets["Received from"].configure(state="normal" if receiving else "disabled")
            input_widgets["Project/site"].configure(state="normal")
            if receiving:
                pass
            elif issuance:
                supplier_var.set("")
            else:
                supplier_var.set("")

        def update_material_fields(*_):
            existing = bool(material_tree.selection())
            input_widgets["New item name"].configure(state="disabled" if existing else "normal")
            input_widgets["Location"].configure(state="disabled" if existing else "normal")
            input_widgets["Short description"].configure(state="disabled" if existing else "normal")
            if existing:
                name_var.set("")
                location_var.set("")
                description_var.set("")

        type_var.trace_add("write", update_type_fields)
        material_tree.bind("<<TreeviewSelect>>", update_material_fields)
        update_type_fields()
        ttk.Label(right, text="Select an existing material, or fill New item name to create one.", style="Muted.TLabel", wraplength=300).grid(row=12, column=0, columnspan=2, sticky="w", pady=(10, 12))

        def save_entry():
            chosen = material_tree.selection()
            try:
                if chosen:
                    material_id = int(chosen[0])
                else:
                    if not name_var.get().strip(): raise ValueError("Select an existing material or enter a new item name.")
                    material_id = create_material(self.db_path, name_var.get(), location_var.get() or None)
                    if description_var.get().strip():
                        set_material_description(self.db_path, material_id, description_var.get())
                if not date_var.get().strip(): raise ValueError("Date is required.")
                quantity = float(qty_var.get())
                if quantity <= 0: raise ValueError("Quantity must be greater than zero.")
                register_project(self.db_path, project_var.get())
                add_paper_record(self.db_path, material_id, date_var.get().strip(), type_var.get(), quantity, ref_var.get().strip() or None, notes_var.get().strip() or None, unit_var.get().strip() or None, supplier_var.get().strip() or None, project_var.get().strip() or None)
            except Exception as exc:
                messagebox.showerror("Could not save entry", str(exc), parent=dialog); return
            dialog.destroy(); self.show_all(); self.status.configure(text="Saved successfully • Stock-card entry recorded")
        ttk.Button(right, text="Batch entry for project/site", command=lambda: self.open_document_window(batch_mode=True)).grid(row=13, column=0, columnspan=2, sticky="ew", pady=(4, 0))
        ttk.Button(right, text="Confirm and save entry", style="Accent.TButton", command=save_entry).grid(row=14, column=0, columnspan=2, sticky="ew", pady=(4, 0))

    def open_batch_window(self):
        # Legacy compatibility path; multi-line forms now use New Warehouse Document.
        dialog = tk.Toplevel(self); dialog.title("Batch entry by project site"); dialog.geometry("900x760"); dialog.minsize(800, 620); dialog.transient(self); dialog.grab_set()
        panel = ttk.Frame(dialog, style="Panel.TFrame", padding=16); panel.pack(fill="both", expand=True)
        ttk.Label(panel, text="Batch entry by project site", style="Title.TLabel").pack(anchor="w")
        ttk.Label(panel, text="Enter the shared form details once, then add each material line below.", style="Muted.TLabel").pack(anchor="w", pady=(2, 12))
        head = ttk.Frame(panel, style="Panel.TFrame"); head.pack(fill="x")
        date_v, type_v, ref_v, site_v, from_v, to_v = (tk.StringVar() for _ in range(6)); type_v.set("ISSUANCE")
        for label, var, values in [("Date",date_v,None),("Type",type_v,["RECEIVING","ISSUANCE","PULL_OUT","OPENING_BALANCE","PHYSICAL_COUNT"]),("Form/reference",ref_v,None),("Project/site",site_v,None),("Received from",from_v,None),("Issued to",to_v,None)]:
            cell=ttk.Frame(head, style="Panel.TFrame"); cell.pack(side="left", padx=(0,10)); ttk.Label(cell,text=label,style="Muted.TLabel").pack(anchor="w")
            w=ttk.Combobox(cell,textvariable=var,values=values,state="readonly",width=18) if values else (DateEntry(cell,textvariable=var,width=16,date_pattern="yyyy-mm-dd") if label == "Date" else ttk.Entry(cell,textvariable=var,width=18)); w.pack()
            if label in {"Project/site", "Issued to"}: self.attach_project_autocomplete(w, var, dialog)
        ttk.Label(panel,text="Material lines (type the material name; matching existing material is used)",style="Muted.TLabel").pack(anchor="w",pady=(16,4))
        grid=ttk.Frame(panel,style="Panel.TFrame"); grid.pack(fill="both",expand=True)
        ttk.Label(grid,text="Material").grid(row=0,column=0,sticky="w"); ttk.Label(grid,text="Quantity").grid(row=0,column=1,sticky="w"); ttk.Label(grid,text="Unit").grid(row=0,column=2,sticky="w")
        rows=[]
        for i in range(12):
            nv,qv,uv=tk.StringVar(),tk.StringVar(),tk.StringVar(); rows.append((nv,qv,uv)); ttk.Entry(grid,textvariable=nv,width=55).grid(row=i+1,column=0,padx=(0,8),pady=3); ttk.Entry(grid,textvariable=qv,width=14).grid(row=i+1,column=1,padx=(0,8),pady=3); ttk.Entry(grid,textvariable=uv,width=14).grid(row=i+1,column=2,pady=3)
        def save():
            conn=connect(self.db_path); lines=[]
            try:
                for nv,qv,uv in rows:
                    if not nv.get().strip(): continue
                    found=conn.execute("SELECT id FROM materials WHERE normalized_name=?", (nv.get().strip().upper(),)).fetchone()
                    if not found: raise ValueError(f"Material not found: {nv.get()}. Select an existing material name.")
                    lines.append({"material_id":found["id"],"quantity":float(qv.get()),"unit":uv.get().strip() or None})
                register_project(self.db_path, site_v.get()); register_project(self.db_path, to_v.get())
                add_paper_batch(self.db_path,date_v.get().strip(),type_v.get(),ref_v.get().strip() or None,site_v.get().strip() or None,lines,from_v.get().strip() or None,to_v.get().strip() or None)
            except Exception as exc: messagebox.showerror("Could not save batch",str(exc),parent=dialog); return
            finally: conn.close()
            dialog.destroy(); self.show_all(); self.status.configure(text=f"Saved successfully • {len(lines)} stock-card entries recorded")
        ttk.Button(panel,text="Confirm and save batch entries",style="Accent.TButton",command=save).pack(anchor="e",pady=(12,0))

    def open_document_window(self, batch_mode=False, document_type="ISSUANCE"):
        dialog=tk.Toplevel(self); dialog.title("Batch stock card entry" if batch_mode else ("Receiving report" if document_type == "RECEIVING" else "Issuance slip")); dialog.geometry("980x820"); dialog.minsize(900,720); dialog.transient(self); dialog.grab_set()
        panel=ttk.Frame(dialog,style="Panel.TFrame",padding=16); panel.pack(fill="both",expand=True)
        ttk.Label(panel,text="Batch stock card entry" if batch_mode else ("Receiving report" if document_type == "RECEIVING" else "Issuance slip"),style="Title.TLabel").pack(anchor="w")
        top=ttk.Frame(panel,style="Panel.TFrame"); top.pack(fill="x",pady=(12,8))
        vals={k:tk.StringVar() for k in ["date","number","site","po","supplier","dr","vehicle","driver","received","released","issues","remarks"]}; typ=tk.StringVar(value=document_type)
        type_options = ["ISSUANCE", "RECEIVING"] if batch_mode else [document_type]
        fields=[("Type",typ,type_options),("Date / date received",vals["date"],None),("PO / RR / IS number",vals["number"],None),("Project/site",vals["site"],None),("PO no.",vals["po"],None),("Supplier / external provider",vals["supplier"],None),("DR/Invoice no.",vals["dr"],None),("Vehicle",vals["vehicle"],None),("Driver",vals["driver"],None),("Received by",vals["received"],None),("Released by",vals["released"],None),("Issues",vals["issues"],None),("Remarks",vals["remarks"],None)]
        for i,(lab,var,options) in enumerate(fields):
            if lab in {"Issues", "Remarks"}:
                grid_row, grid_col = 5 + (0 if lab == "Issues" else 1), 0
            else:
                grid_row, grid_col = i//3, i%3
            cell=ttk.Frame(top,style="Panel.TFrame"); cell.grid(row=grid_row,column=grid_col,columnspan=3 if lab in {"Issues", "Remarks"} else 1,padx=5,pady=4,sticky="w"); ttk.Label(cell,text=lab,style="Muted.TLabel").pack(anchor="w"); field_width = 92 if lab in {"Issues", "Remarks"} else 24; w=ttk.Combobox(cell,textvariable=var,values=options,state="readonly",width=22) if options else (DateEntry(cell,textvariable=var,width=20,date_pattern="yyyy-mm-dd") if lab.startswith("Date") else ttk.Entry(cell,textvariable=var,width=field_width)); w.pack(anchor="w")
            if lab == "Project/site": self.attach_project_autocomplete(w, var, dialog)
            if lab == "Supplier / external provider": self.attach_supplier_autocomplete(w, var, dialog)
            if lab == "Vehicle": self.attach_recurring_autocomplete(w, var, dialog, ["vehicle"])
            if lab == "Driver": self.attach_recurring_autocomplete(w, var, dialog, ["driver"])
            if lab == "Received by": self.attach_recurring_autocomplete(w, var, dialog, ["received_by"])
            if lab == "Released by": self.attach_recurring_autocomplete(w, var, dialog, ["released_by"])
            if lab == "PO / RR / IS number": self.attach_recurring_autocomplete(w, var, dialog, ["document_number", "reference_number", "po_number"])
            if lab == "PO no.": self.attach_recurring_autocomplete(w, var, dialog, ["po_number"])
        ttk.Label(panel,text="Stock card lines" if batch_mode else "Material lines",style="Title.TLabel").pack(anchor="w",pady=(8,4))
        line_holder=ttk.Frame(panel,style="Panel.TFrame"); line_holder.pack(fill="both",expand=True)
        line_canvas=tk.Canvas(line_holder, height=280, highlightthickness=0, background="white")
        line_scroll=ttk.Scrollbar(line_holder, orient="vertical", command=line_canvas.yview); line_scroll.pack(side="right",fill="y"); line_canvas.pack(side="left",fill="both",expand=True); line_canvas.configure(yscrollcommand=line_scroll.set)
        grid=ttk.Frame(line_canvas,style="Panel.TFrame"); grid_window=line_canvas.create_window((0,0),window=grid,anchor="nw")
        grid.bind("<Configure>", lambda _e: line_canvas.configure(scrollregion=line_canvas.bbox("all")))
        line_canvas.bind("<Configure>", lambda e: line_canvas.itemconfigure(grid_window, width=e.width))
        for c,t in enumerate(("Material name","Quantity","Unit")): ttk.Label(grid,text=t).grid(row=0,column=c,sticky="w",padx=4)
        material_names = [r["name"] for r in connect(self.db_path).execute("SELECT name FROM materials ORDER BY name LIMIT 3000").fetchall()]
        lines=[]
        for i in range(12):
            a,b,c=(tk.StringVar() for _ in range(3)); lines.append((a,b,c));
            material_entry = ttk.Entry(grid, textvariable=a, width=55); material_entry.grid(row=i+1, column=0, padx=4, pady=3); self.attach_material_autocomplete(material_entry, a, dialog)
            for col, (var, width) in enumerate(zip((b, c), (12, 12)), 1):
                ttk.Entry(grid, textvariable=var, width=width).grid(row=i+1, column=col, padx=4, pady=3)
        def save():
            conn=connect(self.db_path); out=[]
            try:
                for name,q,u in lines:
                    if not name.get().strip(): continue
                    entered = name.get().strip().upper()
                    row=conn.execute("SELECT id, name FROM materials WHERE normalized_name=? OR UPPER(name)=? OR normalized_name LIKE ? ORDER BY CASE WHEN normalized_name=? THEN 0 ELSE 1 END, name LIMIT 1", (entered, entered, f"%{entered}%", entered)).fetchone()
                    if not row:
                        material_id = create_material(self.db_path, name.get().strip())
                        row = {"id": material_id, "name": name.get().strip()}
                    out.append({"material_id":row["id"],"name":row["name"],"quantity":float(q.get()),"unit":u.get(),"unit_price":None})
                register_project(self.db_path, vals["site"].get())
                if batch_mode:
                    batch_lines=[{"material_id": line["material_id"], "quantity": line["quantity"], "unit": line.get("unit")} for line in out]
                    add_paper_batch(self.db_path, vals["date"].get().strip(), typ.get(), vals["number"].get().strip() or None, vals["site"].get().strip() or None, batch_lines, vals["supplier"].get().strip() or None, vals["site"].get().strip() or None, vals["issues"].get().strip() or None, vals["remarks"].get().strip() or None)
                    dialog.destroy(); self.show_all(); self.status.configure(text=f"Saved successfully • {len(batch_lines)} stock-card entries recorded"); return
                header={"date":vals["date"].get(),"document_number":vals["number"].get(),"project_site":vals["site"].get(),"po_number":vals["po"].get(),"supplier":vals["supplier"].get(),"dr_invoice_number":vals["dr"].get(),"vehicle":vals["vehicle"].get(),"driver":vals["driver"].get(),"received_by":vals["received"].get(),"released_by":vals["released"].get(),"issues":vals["issues"].get(),"remarks":vals["remarks"].get()}
                document_id = create_document(self.db_path,typ.get(),vals["date"].get(),header,out)
                post_document(self.db_path, document_id)
                export_folder = Path(self.settings["export_folder"]) if self.settings.get("export_folder") else Path(__file__).resolve().parent / "05_Exports"
                export_document(export_folder / f"{typ.get()}_{vals['number'].get() or 'draft'}.xlsx", typ.get(), header, out)
            except Exception as exc: messagebox.showerror("Could not save document",str(exc),parent=dialog); return
            finally: conn.close()
            dialog.destroy(); self.status.configure(text="Saved successfully • Document posted")
        def preview():
            win=tk.Toplevel(dialog); win.title("Document preview"); win.geometry("720x560"); text=tk.Text(win,wrap="word",font=("Segoe UI",10)); text.pack(fill="both",expand=True,padx=12,pady=12)
            text.insert("end", f"{typ.get()}\nDate: {vals['date'].get()}\nDocument no.: {vals['number'].get()}\nProject/site: {vals['site'].get()}\nPO no.: {vals['po'].get()}\nSupplier: {vals['supplier'].get()}\nDR/Invoice: {vals['dr'].get()}\n\nMATERIAL LINES\n")
            for name,q,u in lines:
                if name.get().strip(): text.insert("end", f"{name.get()} | Qty: {q.get()} | Unit: {u.get()}\n")
            text.configure(state="disabled")
        footer=ttk.Frame(panel,style="Panel.TFrame"); footer.pack(fill="x",side="bottom",pady=(10,0))
        ttk.Button(footer,text="Preview",command=preview).pack(side="right",padx=(8,0))
        ttk.Button(footer,text="Confirm and save batch entries" if batch_mode else "Confirm and save document",style="Accent.TButton",command=save).pack(side="right")
    def add_paper_entry(self):
        if self.selected_id is None:
            messagebox.showwarning("Select a material", "Select a material before adding a paper transaction.")
            return
        try:
            quantity = float(self.qty_var.get())
            add_paper(self.db_path, self.selected_id, self.date_var.get().strip(), self.type_var.get(), quantity, self.ref_var.get().strip() or None, self.notes_var.get().strip() or None)
        except Exception as exc:
            messagebox.showerror("Could not add entry", str(exc))
            return
        self.date_var.set("")
        self.qty_var.set("")
        self.ref_var.set("")
        self.notes_var.set("")
        self.select_material()
        self.status.configure(text="Saved successfully • Stock-card transaction recorded")


if __name__ == "__main__":
    StockCardApp().mainloop()
