from pathlib import Path
from openpyxl import load_workbook


def export_document(output_path: Path, document_type: str, header: dict, lines: list[dict]) -> Path:
    template_dir = Path(__file__).resolve().parent / "TemplateForms"
    receiving = document_type.upper() == "RECEIVING"
    template = template_dir / ("RCVNG EXCEL.xlsx" if receiving else "ISSUANCE EXCEL.xlsx")
    if not template.exists():
        raise FileNotFoundError(f"Template not found: {template}")
    wb = load_workbook(template)
    ws = wb.active
    if receiving:
        ws["L6"] = header.get("document_number"); ws["B6"] = header.get("supplier"); ws["B7"] = header.get("date"); ws["B8"] = header.get("project_site")
        ws["L7"] = header.get("dr_invoice_number"); ws["L8"] = header.get("po_number"); ws["E30"] = header.get("received_by")
        for index, line in enumerate(lines[:13], 15):
            ws[f"B{index}"] = line.get("quantity"); ws[f"C{index}"] = line.get("unit"); ws[f"F{index}"] = line.get("name"); ws[f"H{index}"] = line.get("unit_price"); ws[f"K{index}"] = f"=B{index}*H{index}"
    else:
        ws["G5"] = header.get("document_number"); ws["B5"] = header.get("project_site"); ws["B6"] = header.get("vehicle"); ws["B7"] = header.get("driver"); ws["G6"] = header.get("date")
        for index, line in enumerate(lines[:12], 11):
            ws[f"A{index}"] = header.get("po_number"); ws[f"B{index}"] = line.get("name"); ws[f"F{index}"] = line.get("quantity"); ws[f"G{index}"] = line.get("unit")
        ws["B30"] = header.get("received_by"); ws["F30"] = header.get("released_by")
    output_path = output_path.with_suffix(".xlsx"); output_path.parent.mkdir(parents=True, exist_ok=True); wb.save(output_path); return output_path
