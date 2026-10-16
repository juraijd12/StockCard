from __future__ import annotations

import json
import math
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl


# ============================================================================
# CONFIGURATION
# ============================================================================

INPUT_FILE = str(Path(__file__).resolve().parent / "CEILING WORKS.xlsx")


# ============================================================================
# BASIC HELPERS
# ============================================================================

def clean_text(value: Any) -> Optional[str]:
    if value is None:
        return None

    text = str(value).strip()

    if not text:
        return None

    return text


def normalize_text(value: Any) -> str:
    text = clean_text(value)

    if text is None:
        return ""

    return re.sub(r"\s+", " ", text).strip().upper()


def safe_json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()

    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None

    return value


# ============================================================================
# DATE HANDLING
# ============================================================================

def parse_date(value: Any) -> Optional[date]:

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel

            converted = from_excel(value)

            if isinstance(converted, datetime):
                return converted.date()

            if isinstance(converted, date):
                return converted

        except Exception:
            pass

    text = clean_text(value)

    if not text:
        return None

    formats = [
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%m/%d/%Y",
        "%m/%d/%y",
        "%d/%m/%Y",
        "%d/%m/%y",
        "%B %d, %Y",
        "%b %d, %Y",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    return None


def format_date(value: Any) -> Optional[str]:

    parsed = parse_date(value)

    if parsed is None:
        return clean_text(value)

    return parsed.isoformat()


# ============================================================================
# QUANTITY HANDLING
# ============================================================================

def parse_quantity(value: Any) -> dict[str, Any]:

    if value is None:
        return {
            "raw": None,
            "value": None,
            "unit": None,
        }

    if isinstance(value, bool):
        return {
            "raw": str(value),
            "value": None,
            "unit": None,
        }

    if isinstance(value, (int, float)):

        numeric = float(value)

        if numeric.is_integer():
            numeric = int(numeric)

        return {
            "raw": str(value),
            "value": numeric,
            "unit": None,
        }

    text = clean_text(value)

    if not text:
        return {
            "raw": None,
            "value": None,
            "unit": None,
        }

    match = re.match(
        r"^\s*([-+]?\d+(?:\.\d+)?)\s*([A-Za-z]+)?\s*$",
        text,
    )

    if match:

        try:
            numeric = float(match.group(1))

            if numeric.is_integer():
                numeric = int(numeric)

            unit = match.group(2)

            return {
                "raw": text,
                "value": numeric,
                "unit": unit.upper() if unit else None,
            }

        except ValueError:
            pass

    match = re.match(
        r"^\s*([-+]?\d+(?:\.\d+)?)\s*(.*?)\s*$",
        text,
    )

    if match:

        try:
            numeric = float(match.group(1))

            if numeric.is_integer():
                numeric = int(numeric)

            remainder = clean_text(match.group(2))

            return {
                "raw": text,
                "value": numeric,
                "unit": remainder.upper() if remainder else None,
            }

        except ValueError:
            pass

    return {
        "raw": text,
        "value": None,
        "unit": None,
    }


# ============================================================================
# STOCK CARD BLOCK DETECTION
# ============================================================================

def find_stock_card_rows(ws) -> list[int]:

    rows = []

    for row in range(1, ws.max_row + 1):

        for col in range(1, ws.max_column + 1):

            if normalize_text(ws.cell(row, col).value) == "STOCK CARD":
                rows.append(row)
                break

    return rows


# ============================================================================
# HEADER DETECTION
# ============================================================================

def find_header_row(
    ws,
    block_start_row: int,
    block_end_row: int,
) -> Optional[int]:

    required = {
        "DATE",
        "RECEIVING",
        "ISSUANCE",
        "QUANTITY",
        "BALANCE ON HAND",
    }

    for row in range(block_start_row, block_end_row + 1):

        values = {
            normalize_text(ws.cell(row, col).value)
            for col in range(1, ws.max_column + 1)
        }

        if required.issubset(values):
            return row

    return None


# ============================================================================
# SIDE DETECTION
# ============================================================================

def detect_card_sides(ws, header_row: int) -> list[dict[str, Any]]:

    sides = []

    # A:H
    if (
        normalize_text(ws.cell(header_row, 1).value) == "DATE"
        and
        normalize_text(ws.cell(header_row, 8).value)
        == "BALANCE ON HAND"
    ):
        sides.append(
            {
                "side_number": 1,
                "start_col": 1,
                "end_col": 8,
                "columns": "A:H",
            }
        )

    # J:Q
    if (
        normalize_text(ws.cell(header_row, 10).value) == "DATE"
        and
        normalize_text(ws.cell(header_row, 17).value)
        == "BALANCE ON HAND"
    ):
        sides.append(
            {
                "side_number": 2,
                "start_col": 10,
                "end_col": 17,
                "columns": "J:Q",
            }
        )

    return sides


# ============================================================================
# ITEM / LOCATION EXTRACTION
# ============================================================================

def extract_item_location(
    ws,
    side_start_col: int,
    header_row: int,
) -> tuple[Optional[str], Optional[str]]:
    """
    Exact workbook layout:

        header_row - 4 : REORDER POINTS
        header_row - 3 : ITEM
        header_row - 2 : LOCATION
        header_row - 1 : blank
        header_row     : DATE

    Left side:

        A = label
        B = value

    Right side:

        J = label
        K = value
    """

    # THIS IS THE IMPORTANT FIX.
    item_row = header_row - 3
    location_row = header_row - 2

    item = None
    location = None

    item_label = normalize_text(
        ws.cell(item_row, side_start_col).value
    )

    location_label = normalize_text(
        ws.cell(location_row, side_start_col).value
    )

    if item_label == "ITEM:":
        item = clean_text(
            ws.cell(
                item_row,
                side_start_col + 1
            ).value
        )

    if location_label == "LOCATION:":
        location = clean_text(
            ws.cell(
                location_row,
                side_start_col + 1
            ).value
        )

    return item, location


# ============================================================================
# COLUMN MAPPING
# ============================================================================

def build_column_mapping(
    side_start_col: int
) -> dict[str, str]:

    if side_start_col == 1:

        return {
            "date": "A",
            "receiving_ref": "B",
            "provider": "C",
            "issuance_ref": "D",
            "issued_to": "E",
            "qty_in": "F",
            "qty_out": "G",
            "balance": "H",
        }

    if side_start_col == 10:

        return {
            "date": "J",
            "receiving_ref": "K",
            "provider": "L",
            "issuance_ref": "M",
            "issued_to": "N",
            "qty_in": "O",
            "qty_out": "P",
            "balance": "Q",
        }

    raise ValueError(
        f"Unsupported stock-card side starting column: "
        f"{side_start_col}"
    )


# ============================================================================
# TRANSACTION EXTRACTION
# ============================================================================

def extract_transactions(
    ws,
    header_row: int,
    block_end_row: int,
    column_mapping: dict[str, str],
) -> list[dict[str, Any]]:

    from openpyxl.utils.cell import column_index_from_string

    transactions = []

    date_col = column_index_from_string(
        column_mapping["date"]
    )

    receiving_ref_col = column_index_from_string(
        column_mapping["receiving_ref"]
    )

    provider_col = column_index_from_string(
        column_mapping["provider"]
    )

    issuance_ref_col = column_index_from_string(
        column_mapping["issuance_ref"]
    )

    issued_to_col = column_index_from_string(
        column_mapping["issued_to"]
    )

    qty_in_col = column_index_from_string(
        column_mapping["qty_in"]
    )

    qty_out_col = column_index_from_string(
        column_mapping["qty_out"]
    )

    balance_col = column_index_from_string(
        column_mapping["balance"]
    )

    # Detailed header is header_row + 1.
    # Actual transaction data begins header_row + 2.
    data_start_row = header_row + 2

    for row in range(
        data_start_row,
        block_end_row + 1
    ):

        date_raw = ws.cell(
            row,
            date_col
        ).value

        parsed_date = parse_date(date_raw)

        if parsed_date is None:
            continue

        receiving_ref = clean_text(
            ws.cell(
                row,
                receiving_ref_col
            ).value
        )

        provider = clean_text(
            ws.cell(
                row,
                provider_col
            ).value
        )

        issuance_ref = clean_text(
            ws.cell(
                row,
                issuance_ref_col
            ).value
        )

        issued_to = clean_text(
            ws.cell(
                row,
                issued_to_col
            ).value
        )

        qty_in_raw = ws.cell(
            row,
            qty_in_col
        ).value

        qty_out_raw = ws.cell(
            row,
            qty_out_col
        ).value

        balance_raw = ws.cell(
            row,
            balance_col
        ).value

        qty_in = parse_quantity(
            qty_in_raw
        )

        qty_out = parse_quantity(
            qty_out_raw
        )

        balance = parse_quantity(
            balance_raw
        )

        # Determine transaction type.

        if receiving_ref and not issuance_ref:

            transaction_type = "RECEIVING"
            reference = receiving_ref

        elif issuance_ref and not receiving_ref:

            transaction_type = "ISSUANCE"
            reference = issuance_ref

        elif receiving_ref and issuance_ref:

            if (
                qty_in["value"] is not None
                and
                qty_out["value"] is None
            ):

                transaction_type = "RECEIVING"
                reference = receiving_ref

            elif (
                qty_out["value"] is not None
                and
                qty_in["value"] is None
            ):

                transaction_type = "ISSUANCE"
                reference = issuance_ref

            else:

                transaction_type = "MIXED"
                reference = receiving_ref

        else:

            if (
                qty_in["value"] is not None
                and
                qty_out["value"] is None
            ):

                transaction_type = "RECEIVING"

            elif (
                qty_out["value"] is not None
                and
                qty_in["value"] is None
            ):

                transaction_type = "ISSUANCE"

            else:

                transaction_type = "UNKNOWN"

            reference = None

        transactions.append(
            {
                "source_row": row,

                "date": format_date(
                    date_raw
                ),

                "type": transaction_type,

                "reference": reference,

                "receiving_ref": receiving_ref,

                "provider": provider,

                "issuance_ref": issuance_ref,

                "issued_to": issued_to,

                "quantity_in": qty_in,

                "quantity_out": qty_out,

                "recorded_balance": balance,

                "raw": {
                    "date": safe_json_value(
                        date_raw
                    ),

                    "receiving_ref": safe_json_value(
                        ws.cell(
                            row,
                            receiving_ref_col
                        ).value
                    ),

                    "provider": safe_json_value(
                        ws.cell(
                            row,
                            provider_col
                        ).value
                    ),

                    "issuance_ref": safe_json_value(
                        ws.cell(
                            row,
                            issuance_ref_col
                        ).value
                    ),

                    "issued_to": safe_json_value(
                        ws.cell(
                            row,
                            issued_to_col
                        ).value
                    ),

                    "quantity_in": safe_json_value(
                        qty_in_raw
                    ),

                    "quantity_out": safe_json_value(
                        qty_out_raw
                    ),

                    "balance": safe_json_value(
                        balance_raw
                    ),
                },
            }
        )

    return transactions


# ============================================================================
# EXTRACTION VALIDATION
# ============================================================================

def validate_transactions(
    transactions: list[dict[str, Any]]
) -> list[str]:

    warnings = []

    if not transactions:
        return warnings

    previous_date = None

    chronology_warning_added = False

    for transaction in transactions:

        current_date = parse_date(
            transaction.get("date")
        )

        if current_date is None:
            continue

        if (
            previous_date is not None
            and
            current_date < previous_date
        ):

            if not chronology_warning_added:

                warnings.append(
                    "Transactions are not in chronological order."
                )

                chronology_warning_added = True

        previous_date = current_date

    for transaction in transactions:

        row = transaction["source_row"]

        for field_name in (
            "quantity_in",
            "quantity_out",
        ):

            quantity = transaction[field_name]

            if quantity.get("unit"):

                warnings.append(
                    f"Row {row}: quantity contains unit: "
                    f"{quantity.get('raw')}"
                )

    return warnings


# ============================================================================
# SIDE EXTRACTION
# ============================================================================

def extract_side(
    ws,
    side_info: dict[str, Any],
    header_row: int,
    block_end_row: int,
) -> dict[str, Any]:

    side_start_col = side_info["start_col"]

    item, location = extract_item_location(
        ws,
        side_start_col,
        header_row,
    )

    column_mapping = build_column_mapping(
        side_start_col
    )

    transactions = extract_transactions(
        ws,
        header_row,
        block_end_row,
        column_mapping,
    )

    warnings = validate_transactions(
        transactions
    )

    return {
        "side_number": side_info["side_number"],
        "columns": side_info["columns"],

        "item": item,
        "location": location,

        "column_mapping": column_mapping,

        "transactions": transactions,

        "warnings": warnings,
    }


# ============================================================================
# CONSOLE DISPLAY
# ============================================================================

def display_side(
    side: dict[str, Any]
) -> None:

    print()
    print(
        f"    SIDE #{side['side_number']}"
    )

    print(
        f"    Columns: {side['columns']}"
    )

    print(
        "    Item: "
        +
        (
            side["item"]
            if side["item"] is not None
            else "[BLANK]"
        )
    )

    print(
        "    Location: "
        +
        (
            side["location"]
            if side["location"] is not None
            else "[BLANK]"
        )
    )

    transactions = side["transactions"]

    print(
        f"    Transactions: {len(transactions)}"
    )

    if transactions:

        print()
        print(
            "      TRANSACTIONS:"
        )

        for transaction in transactions:

            if transaction["type"] == "ISSUANCE":
                quantity = transaction[
                    "quantity_out"
                ]

            else:
                quantity = transaction[
                    "quantity_in"
                ]

            print(
                f"        Row "
                f"{transaction['source_row']}: "
                f"{transaction['date']} | "
                f"{transaction['type']} | "
                f"{transaction['reference']} | "
                f"{quantity.get('raw')}"
            )

    if side["warnings"]:

        print()
        print(
            "      EXTRACTION WARNINGS:"
        )

        for warning in side["warnings"]:

            print(
                f"        WARNING: {warning}"
            )

    if not transactions:

        print()
        print(
            "      INFO: No transactions detected."
        )


# ============================================================================
# WORKBOOK INSPECTION
# ============================================================================

def inspect_workbook(
    input_file: str,
    write_json: bool = True,
    output_dir: Optional[str] = None,
) -> dict[str, Any]:

    input_path = Path(input_file)

    if not input_path.exists():

        raise FileNotFoundError(
            f"Input file not found: "
            f"{input_path}"
        )

    print(
        "=" * 70
    )

    print(
        f"WORKBOOK: {input_path.name}"
    )

    print(
        "=" * 70
    )

    wb = openpyxl.load_workbook(
        input_path,
        data_only=False,
    )

    report = {
        "schema_version": "2.2",

        "workbook": input_path.name,

        "source_file": str(
            input_path
        ),

        "sheets": [],

        "summary": {
            "stock_card_blocks": 0,
            "stock_card_sides": 0,
            "transactions": 0,
            "warnings": 0,
        },
    }

    for ws in wb.worksheets:

        print()

        print(
            f"SHEET: {ws.title}"
        )

        print(
            f"DIMENSIONS: "
            f"{ws.max_row} rows × "
            f"{ws.max_column} columns"
        )

        stock_card_rows = find_stock_card_rows(
            ws
        )

        print(
            f"  STOCK CARD blocks: "
            f"{len(stock_card_rows)}"
        )

        sheet_report = {
            "sheet_name": ws.title,

            "max_row": ws.max_row,

            "max_column": ws.max_column,

            "stock_cards": [],
        }

        for index, block_start_row in enumerate(
            stock_card_rows,
            start=1,
        ):

            if index < len(stock_card_rows):

                block_end_row = (
                    stock_card_rows[index] - 1
                )

            else:

                block_end_row = ws.max_row

            print()

            print(
                f"  STOCK CARD BLOCK #{index}"
            )

            print(
                f"  Block: rows "
                f"{block_start_row}-"
                f"{block_end_row}"
            )

            header_row = find_header_row(
                ws,
                block_start_row,
                block_end_row,
            )

            if header_row is None:

                print(
                    "  WARNING: Could not find "
                    "transaction header."
                )

                sheet_report[
                    "stock_cards"
                ].append(
                    {
                        "block_number": index,

                        "block_start_row":
                            block_start_row,

                        "block_end_row":
                            block_end_row,

                        "header_row": None,

                        "sides": [],

                        "warnings": [
                            "Could not find "
                            "transaction header."
                        ],
                    }
                )

                report[
                    "summary"
                ]["warnings"] += 1

                continue

            print(
                f"  Header: row {header_row}"
            )

            sides = detect_card_sides(
                ws,
                header_row,
            )

            print(
                f"  Sides detected: "
                f"{len(sides)}"
            )

            block_report = {
                "block_number": index,

                "block_start_row":
                    block_start_row,

                "block_end_row":
                    block_end_row,

                "header_row":
                    header_row,

                "sides": [],

                "warnings": [],
            }

            for side_info in sides:

                side = extract_side(
                    ws,
                    side_info,
                    header_row,
                    block_end_row,
                )

                block_report[
                    "sides"
                ].append(side)

                display_side(side)

                report[
                    "summary"
                ]["stock_card_sides"] += 1

                report[
                    "summary"
                ]["transactions"] += len(
                    side["transactions"]
                )

                report[
                    "summary"
                ]["warnings"] += len(
                    side["warnings"]
                )

            sheet_report[
                "stock_cards"
            ].append(
                block_report
            )

            report[
                "summary"
            ]["stock_card_blocks"] += 1

        report[
            "sheets"
        ].append(
            sheet_report
        )

    # ========================================================================
    # SUMMARY
    # ========================================================================

    print()

    print(
        "=" * 70
    )

    print(
        "SUMMARY"
    )

    print(
        "=" * 70
    )

    print(
        f"Stock card blocks detected: "
        f"{report['summary']['stock_card_blocks']}"
    )

    print(
        f"Stock card sides detected: "
        f"{report['summary']['stock_card_sides']}"
    )

    print(
        f"Transactions detected: "
        f"{report['summary']['transactions']}"
    )

    print(
        f"Warnings: "
        f"{report['summary']['warnings']}"
    )

    wb.close()

    if write_json:
        destination = Path(output_dir) if output_dir else input_path.parent
        destination.mkdir(parents=True, exist_ok=True)
        output_file = destination / f"{input_path.stem}_inspection.json"

        with open(output_file, "w", encoding="utf-8") as f:
            json.dump(report, f, indent=2, ensure_ascii=False)

        print()
        print("JSON report saved to:")
        print(output_file)

    return report


# ============================================================================
# MAIN
# ============================================================================

if __name__ == "__main__":

    try:

        input_file = INPUT_FILE

        if len(sys.argv) > 1:

            cli_path = Path(
                sys.argv[1]
            )

            if cli_path.exists():
                input_file = str(
                    cli_path
                )

        inspect_workbook(
            input_file
        )

    except Exception as exc:

        print()

        print(
            "=" * 70
        )

        print(
            "ERROR"
        )

        print(
            "=" * 70
        )

        print(
            str(exc)
        )

        raise   
